"""Build a local review page from the review workbook.

Offline and read-only: it reads the current review workbook (review/legacy_tag_review_*.xlsx)
and writes review/review_page.html, a page you open in your browser. It shows one card per
photo with each tagged object's current tags, status and fix note, and lets you fill in the
same change columns and Review notes as the workbook. Edits are kept in the browser as you
type. "Export edits" downloads them as a small JSON file (the differences from the
workbook, with the value each field had), which is merged into the workbook in a separate,
checked step. The page never writes to the workbook or to OpenLitterMap.

    python3 scripts/build_review_page.py
    python3 scripts/build_review_page.py --workbook review/legacy_tag_review_<date>.xlsx
"""
import argparse
import glob
import json
import os
import sys

import openpyxl

EDIT_FIELDS = ["Change_To_OLMKey", "Change_to_OLM_secondary", "Change_to_OLM_Material",
               "Change_to_OLM_Custom", "Review notes"]

PAGE = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>MROLM Review</title>
<style>
  :root { --bg: #f4f4f2; --panel: #ffffff; --text: #1d1d1b; --muted: #6b6b66; --line: #d9d9d4;
          --accent: #2f6f4f; --edit: #fff4cc; --warn: #b3261e; --ok: #2f6f4f; --review: #8a5a00; --reclass: #6b3fa0; }
  @media (prefers-color-scheme: dark) {
    :root { --bg: #161615; --panel: #22221f; --text: #ececea; --muted: #a3a39c; --line: #3a3a36;
            --accent: #7cc19b; --edit: #4a3f14; --warn: #ff8a80; --ok: #7cc19b; --review: #e0b35a; --reclass: #c6a4f0; }
  }
  * { box-sizing: border-box; }
  body { margin: 0; background: var(--bg); color: var(--text); font: 14px/1.4 system-ui, sans-serif; }
  header { position: sticky; top: 0; z-index: 2; background: var(--panel); border-bottom: 1px solid var(--line);
           padding: 10px 16px; display: flex; flex-wrap: wrap; gap: 8px 14px; align-items: end; }
  h1 { font-size: 16px; margin: 0 8px 0 0; align-self: center; }
  label { display: flex; flex-direction: column; font-size: 12px; color: var(--muted); gap: 2px; }
  label.check { flex-direction: row; align-items: center; gap: 6px; align-self: center; font-size: 13px; color: var(--text); }
  input, select, button { font: inherit; color: var(--text); background: var(--bg); border: 1px solid var(--line);
                          border-radius: 6px; padding: 5px 7px; }
  button { cursor: pointer; } button.primary { background: var(--accent); color: var(--panel); border-color: var(--accent); }
  #ids { width: 170px; } #custom { width: 130px; } #fix { max-width: 240px; } #key { max-width: 220px; }
  .status-line { width: 100%; display: flex; gap: 16px; font-size: 13px; color: var(--muted); }
  .status-line b { color: var(--text); } .status-line .warn { color: var(--warn); }
  main { padding: 16px; display: grid; gap: 14px; grid-template-columns: repeat(auto-fill, minmax(var(--size, 440px), 1fr)); }
  .card { background: var(--panel); border: 1px solid var(--line); border-radius: 8px; overflow: hidden; }
  .photo { position: relative; aspect-ratio: 4 / 3; background: var(--line); }
  .photo a, .photo img { display: block; width: 100%; height: 100%; }
  .photo img { object-fit: contain; background: #111; }
  .cap { position: absolute; top: 0; left: 0; right: 0; padding: 6px 8px; font-weight: 600; color: #fff;
         background: linear-gradient(rgba(0,0,0,.75), rgba(0,0,0,0)); text-shadow: 0 1px 2px #000;
         display: flex; justify-content: space-between; pointer-events: none; }
  .obj { padding: 8px 10px; border-top: 1px solid var(--line); }
  .obj .now { font-weight: 600; }
  .obj .meta { color: var(--muted); font-size: 12px; }
  .badge { display: inline-block; font-size: 11px; font-weight: 700; padding: 0 6px; border-radius: 4px; border: 1px solid; margin-right: 4px; }
  .OK { color: var(--ok); } .REVIEW { color: var(--review); } .RECLASS, .UNCLASS, .UNMAPPED, .ORPHAN { color: var(--reclass); }
  .fields { display: grid; grid-template-columns: 1fr 1fr; gap: 4px 8px; margin-top: 6px; }
  .fields label { font-size: 11px; } .fields input { width: 100%; padding: 4px 6px; }
  .fields .wide { grid-column: 1 / -1; }
  input.changed { background: var(--edit); } input.bad { border-color: var(--warn); outline: 1px solid var(--warn); }
  nav { display: flex; gap: 8px; justify-content: center; align-items: center; padding: 8px 16px 24px; }
  .empty { grid-column: 1 / -1; color: var(--muted); text-align: center; padding: 40px 0; }
</style>
</head>
<body>
<header>
  <h1>MROLM Review</h1>
  <label>Photo ID(s)<input id="ids" placeholder="e.g. 546290, 5529" autocomplete="off"></label>
  <label>From<input id="from" type="date"></label>
  <label>To<input id="to" type="date"></label>
  <label>OLM key<select id="key"><option value="">All keys</option></select></label>
  <label>Status<select id="status"><option value="">All</option></select></label>
  <label>Fix note<select id="fix"><option value="">All</option></select></label>
  <label>Custom tag<input id="custom" placeholder="contains…" autocomplete="off"></label>
  <label class="check"><input id="edited" type="checkbox">Has edits</label>
  <label>Size<input id="size" type="range" min="320" max="800" step="20" value="440"></label>
  <button id="clear" type="button">Clear filters</button>
  <button id="export" type="button" class="primary">Export edits</button>
  <div class="status-line"><span id="count"></span><span id="pending"></span><span id="orphans" class="warn"></span></div>
</header>
<main id="grid"></main>
<nav><button id="prev" type="button">Previous</button><span id="pageinfo"></span><button id="next" type="button">Next</button></nav>
<datalist id="dl-keys"></datalist><datalist id="dl-types"></datalist><datalist id="dl-materials"></datalist>
<script>
const DATA = __DATA__;
const FIELDS = [["Change_To_OLMKey", "New OLM key", "dl-keys"], ["Change_to_OLM_secondary", "New secondary", "dl-types"],
                ["Change_to_OLM_Material", "New material(s)", "dl-materials"], ["Change_to_OLM_Custom", "New custom tag(s)", ""],
                ["Review notes", "Review notes", ""]];
const PER_PAGE = 24;
const STORE = "mrolm-review-edits";
const $ = id => document.getElementById(id);

// Browser edits: {tagId: {field: value}}. Storage can be unavailable; the page still works without it.
let edits = {};
try { edits = JSON.parse(localStorage.getItem(STORE) || "{}"); } catch (e) { edits = {}; }
function save() { try { localStorage.setItem(STORE, JSON.stringify(edits)); } catch (e) {} updateStatus(); }

const objById = {};
DATA.photos.forEach(p => p.objects.forEach(o => { objById[o.tag] = o; o.photo = p; }));
const value = (o, f) => (edits[o.tag] && f in edits[o.tag]) ? edits[o.tag][f] : (o.edit[f] || "");
const changed = (o, f) => value(o, f) !== (o.edit[f] || "");
const hasEdits = o => FIELDS.some(([f]) => value(o, f) !== "");
const keySet = new Set(DATA.keys), typeSet = new Set(DATA.types), matSet = new Set(DATA.materials);
function valid(f, v) {
  if (!v || v === "[remove]") return f !== "Change_To_OLMKey" || v !== "[remove]";
  if (f === "Change_To_OLMKey") return keySet.has(v);
  if (f === "Change_to_OLM_secondary") return typeSet.has(v);
  if (f === "Change_to_OLM_Material") return v.split(";").map(s => s.trim()).filter(Boolean).every(m => matSet.has(m));
  return true;
}

function fill(id, values) { values.forEach(v => $(id).append(new Option(v, v))); }
DATA.keys.forEach(k => $("dl-keys").append(new Option(k)));
DATA.types.forEach(k => $("dl-types").append(new Option(k)));
DATA.materials.forEach(k => $("dl-materials").append(new Option(k)));
const keyCounts = {};
DATA.photos.forEach(p => p.objects.forEach(o => { const k = o.key || "(orphan custom tag)"; keyCounts[k] = (keyCounts[k] || 0) + 1; }));
Object.keys(keyCounts).sort().forEach(k => $("key").add(new Option(`${k} (${keyCounts[k]})`, k)));
fill("status", [...new Set(DATA.photos.flatMap(p => p.objects.map(o => o.status)))].sort());
DATA.fixes.forEach((t, i) => $("fix").add(new Option(t, String(i))));

let page = 0, shown = DATA.photos;
function objectMatches(o) {
  const key = $("key").value, status = $("status").value, fix = $("fix").value, custom = $("custom").value.trim().toLowerCase();
  if (key && (o.key || "(orphan custom tag)") !== key) return false;
  if (status && o.status !== status) return false;
  if (fix && String(o.fix) !== fix) return false;
  if (custom && !o.custom.toLowerCase().includes(custom)) return false;
  if ($("edited").checked && !hasEdits(o)) return false;
  return true;
}
function applyFilters() {
  const ids = $("ids").value.split(/[\s,;]+/).filter(Boolean), from = $("from").value, to = $("to").value;
  shown = DATA.photos.filter(p => {
    const day = p.date.slice(0, 10);
    if (ids.length && !ids.some(x => String(p.id).startsWith(x))) return false;
    if (from && day < from) return false;
    if (to && day > to) return false;
    return p.objects.some(objectMatches);
  });
  page = 0;
  render();
}

function el(tag, cls, text) { const e = document.createElement(tag); if (cls) e.className = cls; if (text != null) e.textContent = text; return e; }
function objectBlock(o) {
  const box = el("div", "obj");
  const now = [o.key || "(orphan custom tag)"];
  if (o.secondary) now.push(o.secondary);
  if (o.material) now.push("material: " + o.material);
  if (o.custom) now.push("custom: " + o.custom);
  now.push("x" + o.qty, o.picked === "yes" ? "picked up" : "NOT picked up");
  box.append(el("div", "now", `#${o.no}  ${now.join(" · ")}`));
  const meta = el("div", "meta");
  const status = o.status.split(" ")[0];
  meta.append(el("span", "badge " + status, o.status), document.createTextNode(o.local || ""));
  if (o.fix !== null) meta.append(el("div", null, "Fix note: " + DATA.fixes[o.fix]));
  if (o.def !== null) meta.title = DATA.defs[o.def];
  box.append(meta);
  const grid = el("div", "fields");
  for (const [f, label, list] of FIELDS) {
    const lab = el("label", f === "Review notes" ? "wide" : null, label);
    const input = el("input");
    input.value = value(o, f);
    if (list) input.setAttribute("list", list);
    const mark = () => { input.classList.toggle("changed", changed(o, f)); input.classList.toggle("bad", !valid(f, input.value.trim())); };
    input.addEventListener("input", () => {
      const v = input.value.trim();
      edits[o.tag] = edits[o.tag] || {};
      if (v === (o.edit[f] || "")) delete edits[o.tag][f]; else edits[o.tag][f] = v;
      if (!Object.keys(edits[o.tag]).length) delete edits[o.tag];
      mark(); save();
    });
    mark();
    lab.append(input);
    grid.append(lab);
  }
  box.append(grid);
  return box;
}
function render() {
  const pages = Math.max(1, Math.ceil(shown.length / PER_PAGE));
  page = Math.min(page, pages - 1);
  const grid = $("grid");
  grid.replaceChildren();
  const slice = shown.slice(page * PER_PAGE, (page + 1) * PER_PAGE);
  if (!slice.length) grid.append(el("p", "empty", "No photos match."));
  for (const p of slice) {
    const card = el("div", "card"), ph = el("div", "photo"), a = el("a"), img = el("img"), cap = el("div", "cap");
    a.href = p.url; a.target = "_blank"; a.rel = "noopener";
    img.loading = "lazy"; img.alt = "OLM photo " + p.id; img.src = p.url;
    a.append(img); cap.append(el("span", null, p.id), el("span", null, p.date)); ph.append(a, cap); card.append(ph);
    p.objects.forEach(o => card.append(objectBlock(o)));
    grid.append(card);
  }
  $("count").textContent = `${shown.length} of ${DATA.photos.length} photos`;
  $("pageinfo").textContent = `Page ${page + 1} of ${pages}`;
  $("prev").disabled = page === 0; $("next").disabled = page >= pages - 1;
  updateStatus();
}
function updateStatus() {
  const live = Object.keys(edits).filter(t => objById[t]);
  const gone = Object.keys(edits).length - live.length;
  $("pending").innerHTML = `<b>${live.length}</b> objects with edits in this browser, not yet in the workbook`;
  $("orphans").textContent = gone ? `${gone} browser edits are for tags no longer in this page (sent or re-tagged); they will not export` : "";
}

$("export").onclick = () => {
  const out = [];
  for (const [tag, fields] of Object.entries(edits)) {
    const o = objById[tag];
    if (!o) continue;
    for (const [f, v] of Object.entries(fields))
      out.push({photo_id: o.photo.id, tag_id: Number(tag), field: f, was: o.edit[f] || "", now: v});
  }
  if (!out.length) { alert("No edits to export."); return; }
  const stamp = new Date().toISOString().slice(0, 19).replace(/[-:T]/g, "");
  const blob = new Blob([JSON.stringify({workbook: DATA.workbook, exported_at: new Date().toISOString(), edits: out}, null, 1)],
                        {type: "application/json"});
  const a = el("a"); a.href = URL.createObjectURL(blob); a.download = `review_edits_${stamp}.json`; a.click();
  URL.revokeObjectURL(a.href);
};
["ids", "from", "to", "key", "status", "fix", "custom", "edited"].forEach(id => $(id).addEventListener("input", applyFilters));
$("size").addEventListener("input", e => document.documentElement.style.setProperty("--size", e.target.value + "px"));
$("prev").onclick = () => { page--; render(); scrollTo(0, 0); };
$("next").onclick = () => { page++; render(); scrollTo(0, 0); };
$("clear").onclick = () => { ["ids", "from", "to", "key", "status", "fix", "custom"].forEach(id => $(id).value = ""); $("edited").checked = false; applyFilters(); };
render();
</script>
</body>
</html>
"""


def newest_workbook():
    files = sorted(f for f in glob.glob(os.path.join("review", "legacy_tag_review_*.xlsx")))
    if not files:
        sys.exit("[CRITICAL ERROR] No review workbook in review/. Run scripts/build_review_workbook.py first.")
    return max(files, key=os.path.getmtime)


def main():
    parser = argparse.ArgumentParser(description="Build the local review page from the review workbook.")
    parser.add_argument("--workbook", help="review workbook (default: the most recently saved one in review/)")
    parser.add_argument("--out", default=os.path.join("review", "review_page.html"))
    args = parser.parse_args()
    workbook = args.workbook or newest_workbook()

    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb["Objects"]
    header = [c.value for c in ws[1]]
    missing = [f for f in ["PhotoID", "OLM_Tag_ID", "PhotoLink", "Date"] + EDIT_FIELDS if f not in header]
    if missing:
        sys.exit(f"[CRITICAL ERROR] {workbook} is missing column(s) {missing}.")

    fixes, defs, photos = [], [], {}

    def index(table, text):
        if not text:
            return None
        if text not in table:
            table.append(text)
        return table.index(text)

    for values in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(header, values))
        if r["PhotoID"] is None:
            continue
        pid = int(r["PhotoID"])
        photo = photos.setdefault(pid, {"id": pid, "date": str(r["Date"] or ""), "url": r["PhotoLink"] or "", "objects": []})
        photo["objects"].append({
            "tag": int(r["OLM_Tag_ID"]), "no": int(r["Object_No"] or 0),
            "key": r["OLM_key"] or "", "secondary": r["OLM_secondary"] or "",
            "material": r["OLM_Material"] or "", "custom": r["OLM_Custom"] or "",
            "qty": r["Quantity"], "picked": r["Picked_Up"] or "",
            "status": r["Status"] or "", "local": r["MROLM local key"] or "",
            "fix": index(fixes, (r["OLM data needs fixes"] or "").strip()),
            "def": index(defs, (r["Notes"] or "").strip()),
            "edit": {f: str(r[f]).strip() for f in EDIT_FIELDS if r[f] not in (None, "")},
        })
    for photo in photos.values():
        photo["objects"].sort(key=lambda o: o["no"])

    lists = wb["Lists"]
    columns = {name: [] for name in ("keys", "types", "materials")}
    for key, typ, mat in lists.iter_rows(min_row=2, max_col=3, values_only=True):
        for name, v in (("keys", key), ("types", typ), ("materials", mat)):
            if v and v != "[remove]":
                columns[name].append(v)

    data = {"workbook": os.path.basename(workbook), "fixes": fixes, "defs": defs,
            "photos": sorted(photos.values(), key=lambda p: p["date"], reverse=True), **columns}
    # "</" is escaped so no text in the data can close the script tag early
    blob = json.dumps(data, separators=(",", ":"), ensure_ascii=False).replace("</", "<\\/")
    with open(args.out, "w", encoding="utf-8") as f:
        f.write(PAGE.replace("__DATA__", blob))
    os.chmod(args.out, 0o600)
    objects = sum(len(p["objects"]) for p in photos.values())
    print(f"[SUCCESS] Wrote {args.out}: {len(photos)} photos, {objects} objects, from {workbook}")


if __name__ == "__main__":
    main()
