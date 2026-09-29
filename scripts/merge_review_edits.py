"""Merge edits exported from the review page into the review workbook.

The review page (scripts/build_review_page.py) exports review_edits_<time>.json: one entry
per changed field, with the value the workbook had when the page was built ("was") and the
new value ("now"). This script matches each entry to its object by PhotoID and OLM_Tag_ID.
It never sends anything to OpenLitterMap.

Without --apply it only reports. With --apply it refuses if the workbook is open in
LibreOffice, saves a snapshot to review/snapshots/, writes the edits, and checks that only
those cells changed. An edit is skipped (and reported) when its tag is no longer in the
workbook, when the workbook value has changed since the page was built (a conflict), or when
a key, type or material is not in OLM's lists. Merge the newest export: each export holds all
of the browser's edits, so older exports are superseded.

    python3 scripts/merge_review_edits.py                       # newest ~/Downloads/review_edits_*.json
    python3 scripts/merge_review_edits.py --edits PATH --apply
"""
import argparse
import glob
import json
import os
import shutil
import sys
from datetime import datetime

import openpyxl

EDIT_FIELDS = ["Change_To_OLMKey", "Change_to_OLM_secondary", "Change_to_OLM_Material",
               "Change_to_OLM_Custom", "Change_to_Picked_Up", "Change_to_Quantity", "Review notes"]
REMOVE = "[remove]"


def text(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return "" if value is None else str(value).strip()


def invalid_reason(field, value, lists):
    """Why a new value can't be used, or None. Blank always clears the cell."""
    if value == "":
        return None
    if field == "Change_To_OLMKey":
        return None if value in lists["keys"] else f"'{value}' is not an OLM key"
    if field == "Change_to_OLM_secondary":
        return None if value in lists["types"] or value == REMOVE else f"'{value}' is not an OLM type"
    if field == "Change_to_Picked_Up":
        return None if value in ("yes", "no") else f"picked up must be yes or no, not '{value}'"
    if field == "Change_to_Quantity":
        ok = value.isdigit() and int(value) >= 1
        return None if ok else f"quantity must be a whole number of 1 or more, not '{value}'"
    if field == "Change_to_OLM_Material":
        if value == REMOVE:
            return None
        unknown = [m.strip() for m in value.split(";") if m.strip() and m.strip() not in lists["materials"]]
        return f"unknown material {unknown}" if unknown else None
    return None


def plan_merge(current, edits, lists):
    """Sort exported edits into outcomes.

    current: {(photo_id, tag_id): {field: workbook value}}.
    Returns {"apply", "already", "conflict", "missing", "invalid", "bad_field"}, each a list of
    (edit, detail) pairs.
    """
    plan = {name: [] for name in ("apply", "already", "conflict", "missing", "invalid", "bad_field")}
    for edit in edits:
        field, was, now = edit["field"], text(edit.get("was")), text(edit.get("now"))
        key = (int(edit["photo_id"]), int(edit["tag_id"]))
        if field not in EDIT_FIELDS:
            plan["bad_field"].append((edit, field))
        elif key not in current:
            plan["missing"].append((edit, "tag not in workbook (sent or re-tagged since the page was built)"))
        elif current[key][field] == now:
            plan["already"].append((edit, now))
        elif current[key][field] != was:
            plan["conflict"].append((edit, f"workbook now has '{current[key][field]}', page expected '{was}'"))
        elif invalid_reason(field, now, lists):
            plan["invalid"].append((edit, invalid_reason(field, now, lists)))
        else:
            plan["apply"].append((edit, now))
    return plan


def read_workbook(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Objects"]
    header = [c.value for c in ws[1]]
    missing = [f for f in ["PhotoID", "OLM_Tag_ID"] + EDIT_FIELDS if f not in header]
    if missing:
        sys.exit(f"[CRITICAL ERROR] {path} is missing column(s) {missing}. Nothing was written.")
    col = {name: header.index(name) + 1 for name in header if name}
    rows, current = {}, {}
    for i in range(2, ws.max_row + 1):
        pid, tag = ws.cell(i, col["PhotoID"]).value, ws.cell(i, col["OLM_Tag_ID"]).value
        if pid is None or tag is None:
            continue
        key = (int(pid), int(tag))
        rows[key] = i
        current[key] = {f: text(ws.cell(i, col[f]).value) for f in EDIT_FIELDS}
    lists = {"keys": set(), "types": set(), "materials": set()}
    for k, t, m in wb["Lists"].iter_rows(min_row=2, max_col=3, values_only=True):
        for name, v in (("keys", k), ("types", t), ("materials", m)):
            if v:
                lists[name].add(v)
    return wb, ws, col, rows, current, lists


def cell_values(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Objects"]
    return [[c.value for c in row] for row in ws.iter_rows()]


def main():
    parser = argparse.ArgumentParser(description="Merge review page edits into the review workbook.")
    parser.add_argument("--edits", help="exported edits file (default: newest review_edits_*.json in ~/Downloads)")
    parser.add_argument("--workbook", help="review workbook (default: the most recently saved one in review/)")
    parser.add_argument("--apply", action="store_true", help="write to the workbook; without it this only reports")
    args = parser.parse_args()

    edits_path = args.edits
    if not edits_path:
        found = sorted(glob.glob(os.path.expanduser("~/Downloads/review_edits_*.json")), key=os.path.getmtime)
        if not found:
            sys.exit("[CRITICAL ERROR] No review_edits_*.json in ~/Downloads. Pass --edits PATH.")
        edits_path = found[-1]
    workbook = args.workbook or max(glob.glob(os.path.join("review", "legacy_tag_review_*.xlsx")), key=os.path.getmtime)

    with open(edits_path, encoding="utf-8") as f:
        exported = json.load(f)
    wb, ws, col, rows, current, lists = read_workbook(workbook)
    plan = plan_merge(current, exported["edits"], lists)

    print(f"[INFO] Edits file: {edits_path} (exported {exported.get('exported_at')}, {len(exported['edits'])} field edits)")
    print(f"[INFO] Workbook:   {workbook}")
    if exported.get("workbook") != os.path.basename(workbook):
        print(f"[WARN] The page was built from {exported.get('workbook')}; matching by tag ID anyway.")
    labels = {"apply": "To apply", "already": "Already in workbook", "conflict": "CONFLICT (skipped)",
              "missing": "Tag not found (skipped)", "invalid": "Invalid value (skipped)", "bad_field": "Unknown field (skipped)"}
    for name, label in labels.items():
        if plan[name]:
            print(f"\n{label}: {len(plan[name])}")
            for edit, detail in plan[name]:
                print(f"    {edit['photo_id']} tag {edit['tag_id']}  {edit['field']}: '{text(edit.get('was'))}' -> "
                      f"'{text(edit.get('now'))}'" + ("" if name in ("apply", "already") else f"  [{detail}]"))

    if not args.apply:
        print("\n[REPORT ONLY] Nothing was written. Add --apply to write the edits listed under 'To apply'.")
        return 0
    if not plan["apply"]:
        print("\n[INFO] Nothing to apply. Nothing was written.")
        return 0
    lock = os.path.join(os.path.dirname(workbook), f".~lock.{os.path.basename(workbook)}#")
    if os.path.exists(lock):
        sys.exit("[CRITICAL ERROR] The workbook is open in LibreOffice. Close it and run again. Nothing was written.")

    snapshots = os.path.join(os.path.dirname(workbook), "snapshots")
    os.makedirs(snapshots, exist_ok=True)
    stem = os.path.splitext(os.path.basename(workbook))[0]
    snapshot = os.path.join(snapshots, f"{stem}_{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    shutil.copy2(workbook, snapshot)

    expected = set()
    for edit, now in plan["apply"]:
        r, c = rows[(int(edit["photo_id"]), int(edit["tag_id"]))], col[edit["field"]]
        # A quantity is stored as a number so the workbook's whole-number check accepts it
        ws.cell(r, c).value = (int(now) if edit["field"] == "Change_to_Quantity" and now else now) or None
        expected.add((r, c))
    wb.save(workbook)

    before, after = cell_values(snapshot), cell_values(workbook)
    changed = {(r + 1, c + 1) for r, row in enumerate(after) for c, v in enumerate(row)
               if r >= len(before) or c >= len(before[r]) or before[r][c] != v}
    if len(before) != len(after) or not changed <= expected:
        print(f"[FAILED] Unexpected cells changed. Restore with: cp '{snapshot}' '{workbook}'")
        return 1
    print(f"\n[SUCCESS] Applied {len(plan['apply'])} edits. Only those cells changed. Snapshot: {snapshot}")
    print("[NEXT] Rebuild the review page (python3 scripts/build_review_page.py) so it shows the merged edits.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
