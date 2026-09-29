"""Dry run of a batch tag replacement in OpenLitterMap (OLM). Sends nothing.

Offline and read-only: it reads your edits in a review workbook (the Change_to_*
columns), the newest raw export in review/raw/ (the backup of every photo's
current tags) and OLM's tag list, then builds the full replacement tag list that
PUT /api/v3/tags would need for each edited photo. It prints before and after,
checks the plan, and writes it to review/batches/ (untracked). There is no send
option here; sending is a separate step that only follows a checked dry run.

    python3 scripts/olm_batch_replace.py --workbook "review/<workbook>.xlsx" --label alcohol-packaging

Request format (from OLM's public code, AddTagsToPhotoAction): per tag,
category_litter_object_id, litter_object_type_id, quantity, picked_up,
materials (IDs), brands ({id, quantity}) and custom_tags (plain text).
"""
import argparse
import glob
import json
import os
import sys
from datetime import datetime

import openpyxl

import crosswalk as cw

REMOVE = "[remove]"
CHANGE_KEY, CHANGE_TYPE, CHANGE_MATERIAL, CHANGE_CUSTOM, CHANGE_PICKED = (
    "Change_To_OLMKey", "Change_to_OLM_secondary", "Change_to_OLM_Material", "Change_to_OLM_Custom",
    "Change_to_Picked_Up")
CHANGE_COLUMNS = (CHANGE_KEY, CHANGE_TYPE, CHANGE_MATERIAL, CHANGE_CUSTOM, CHANGE_PICKED)


def newest(pattern):
    files = sorted(glob.glob(pattern))
    if not files:
        sys.exit(f"[CRITICAL ERROR] No files match {pattern}. Run scripts/export_raw_olm.py first.")
    return files[-1]


class TagList:
    """Lookups from OLM's tag list (review/raw/tags_all_*.json)."""

    def __init__(self, tags_all):
        categories = {c["id"]: c["key"] for c in tags_all["categories"]}
        objects = {o["id"]: o["key"] for o in tags_all["objects"]}
        self.clo_key = {co["id"]: f"{categories.get(co['category_id'])}/{objects.get(co['litter_object_id'])}"
                        for co in tags_all["category_objects"]}
        self.clo_by_key = {key: clo for clo, key in self.clo_key.items()}
        self.type_key = {t["id"]: t["key"] for t in tags_all["types"]}
        self.type_by_key = {key: tid for tid, key in self.type_key.items()}
        self.material_key = {m["id"]: m["key"] for m in tags_all["materials"]}
        self.material_by_key = {key: mid for mid, key in self.material_key.items()}
        self.valid_types = {(x["category_litter_object_id"], x["litter_object_type_id"])
                            for x in tags_all["category_object_types"]}


def read_edits(path):
    """{photo_id: {tag_id: {change column: value}}} for every row with a Change_to_* cell filled."""
    ws = openpyxl.load_workbook(path, data_only=True)["Objects"]
    header = [c.value for c in ws[1]]
    edits = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values))
        change = {name: str(row[name]).strip() for name in CHANGE_COLUMNS if row.get(name) not in (None, "")}
        if change:
            edits.setdefault(int(row["PhotoID"]), {})[int(row["OLM_Tag_ID"])] = change
    return edits


def to_payload(tag):
    """One raw export tag in the PUT /api/v3/tags format, keeping every value as it is.

    A custom tag with no object (an orphan) uses OLM's custom-tag-only format: "custom" and
    "key" (the first custom tag), with no category_litter_object_id.
    """
    extras = tag.get("extra_tags") or []
    payload = {
        "category_litter_object_id": tag.get("category_litter_object_id"),
        "litter_object_type_id": tag.get("litter_object_type_id"),
        "quantity": tag.get("quantity"),
        "picked_up": tag.get("picked_up"),
        "materials": [e["tag"]["id"] for e in extras if e["type"] == "material"],
        "brands": [{"id": e["tag"]["id"], "quantity": e.get("quantity", 1)} for e in extras if e["type"] == "brand"],
        "custom_tags": [e["tag"]["key"] for e in extras if e["type"] == "custom_tag"],
    }
    return as_orphan(payload) if payload["category_litter_object_id"] is None else payload


def is_orphan(payload):
    return "category_litter_object_id" not in payload


def as_orphan(payload):
    """An object-less payload in OLM's custom-tag-only format."""
    customs = payload["custom_tags"]
    return {"custom": True, "key": customs[0] if customs else "", "quantity": payload["quantity"],
            "picked_up": payload["picked_up"], "materials": payload["materials"], "brands": payload["brands"],
            "custom_tags": customs[1:]}


def canonical(payload):
    """Either payload format as one dict with the object fields, for comparing and editing."""
    if not is_orphan(payload):
        return payload
    return {"category_litter_object_id": None, "litter_object_type_id": None, "quantity": payload["quantity"],
            "picked_up": payload["picked_up"], "materials": payload["materials"], "brands": payload["brands"],
            "custom_tags": ([payload["key"]] if payload["key"] else []) + payload["custom_tags"]}


def matches_summary(payloads, photo):
    """True when the rebuilt tags match OLM's own summary for the photo, so unchanged tags replay exactly."""
    summary = (photo.get("summary") or {}).get("tags") or []
    custom_ids = {v: int(k) for k, v in ((photo.get("summary") or {}).get("keys", {}).get("custom_tags") or {}).items()}
    rebuilt = [(p["category_litter_object_id"], p["litter_object_type_id"], p["quantity"], p["picked_up"],
                sorted(p["materials"]), sorted((b["id"], b["quantity"]) for b in p["brands"]),
                sorted(custom_ids.get(c) for c in p["custom_tags"])) for p in map(canonical, payloads)]
    def brand_pairs(brands):
        # The summary stores brands as {"<brand id>": quantity} or a list of such dicts
        dicts = [brands] if isinstance(brands, dict) else brands
        return sorted((int(k), v) for d in dicts for k, v in d.items())

    expected = [(s["clo_id"], s["type_id"], s["quantity"], s["picked_up"], sorted(s["materials"]),
                 brand_pairs(s["brands"]), sorted(s["custom_tags"])) for s in summary]
    return sorted(rebuilt, key=str) == sorted(expected, key=str)


def apply_change(payload, change, tags, problems):
    """Apply one row's Change_to_* cells to a payload tag. Blank keeps, [remove] removes, else the new value.

    An orphan given a key becomes an ordinary object; without a key it stays an orphan.
    """
    new = dict(canonical(payload))
    if CHANGE_KEY in change:
        clo = tags.clo_by_key.get(change[CHANGE_KEY])
        if clo is None:
            problems.append(f"unknown OLM key '{change[CHANGE_KEY]}'")
        else:
            new["category_litter_object_id"] = clo
    if CHANGE_TYPE in change:
        value = change[CHANGE_TYPE]
        if value == REMOVE:
            new["litter_object_type_id"] = None
        elif value in tags.type_by_key:
            new["litter_object_type_id"] = tags.type_by_key[value]
        else:
            problems.append(f"unknown secondary (type) '{value}'")
    if CHANGE_MATERIAL in change:
        value = change[CHANGE_MATERIAL]
        wanted = [] if value == REMOVE else [m.strip() for m in value.split(";") if m.strip()]
        unknown = [m for m in wanted if m not in tags.material_by_key]
        if unknown:
            problems.append(f"unknown material {unknown}")
        new["materials"] = [tags.material_by_key[m] for m in wanted if m in tags.material_by_key]
    if CHANGE_CUSTOM in change:
        value = change[CHANGE_CUSTOM]
        new["custom_tags"] = [] if value == REMOVE else [c.strip() for c in value.split(";") if c.strip()]
    if CHANGE_PICKED in change:
        if change[CHANGE_PICKED] in ("yes", "no"):
            new["picked_up"] = change[CHANGE_PICKED] == "yes"
        else:
            problems.append(f"picked up must be yes or no, not '{change[CHANGE_PICKED]}'")
    if new["category_litter_object_id"] is None:
        if new["litter_object_type_id"] or new["materials"]:
            problems.append("an orphan custom tag needs a new OLM key before it can take a type or material")
        if not new["custom_tags"]:
            problems.append("removing an orphan's only custom tag would leave an empty tag; give it a key instead")
        return as_orphan(new)
    type_id = new["litter_object_type_id"]
    if type_id and (new["category_litter_object_id"], type_id) not in tags.valid_types:
        problems.append(f"type '{tags.type_key.get(type_id)}' is not valid for "
                        f"'{tags.clo_key.get(new['category_litter_object_id'])}'")
    return new


def describe(payload, tags):
    payload = canonical(payload)
    key = tags.clo_key.get(payload["category_litter_object_id"], "(orphan custom tag)")
    parts = [key]
    if payload["litter_object_type_id"]:
        parts.append(f"type={tags.type_key.get(payload['litter_object_type_id'])}")
    if payload["materials"]:
        parts.append("materials=" + "; ".join(tags.material_key.get(m, str(m)) for m in payload["materials"]))
    if payload["brands"]:
        parts.append(f"brands={len(payload['brands'])}")
    if payload["custom_tags"]:
        parts.append("custom=" + "; ".join(payload["custom_tags"]))
    picked = {True: "picked up", False: "NOT picked up", None: "picked_up unknown"}[payload["picked_up"]]
    return f"{' '.join(parts)}  x{payload['quantity']}  ({picked})"


def layer_for(payload, tags, crosswalk_rows):
    if is_orphan(payload):
        return "ORPHAN TAG (not on the map)"
    materials = [tags.material_key.get(m, "") for m in payload["materials"]]
    secondary = tags.type_key.get(payload["litter_object_type_id"], "") if payload["litter_object_type_id"] else ""
    matched, _ = cw.match_tag(crosswalk_rows, tags.clo_key.get(payload["category_litter_object_id"], ""),
                              secondary, materials, payload["custom_tags"])
    if matched is None:
        return "UNMAPPED"
    return f"{cw.classify(matched)}: {matched['MROLM local key']}"


def main():
    parser = argparse.ArgumentParser(description="Dry run of a batch OLM tag replacement. Sends nothing.")
    parser.add_argument("--workbook", required=True, help="review workbook with your Change_to_* edits")
    parser.add_argument("--label", required=True, help="short batch name, e.g. alcohol-packaging")
    parser.add_argument("--raw-dir", default=os.path.join("review", "raw"))
    parser.add_argument("--crosswalk", default=os.path.join("config", "crosswalk.csv"))
    parser.add_argument("--out-dir", default=os.path.join("review", "batches"))
    args = parser.parse_args()

    photos_path = newest(os.path.join(args.raw_dir, "photos_*.json"))
    with open(photos_path, encoding="utf-8") as f:
        photos = {p["id"]: p for p in json.load(f)}
    with open(newest(os.path.join(args.raw_dir, "tags_all_*.json")), encoding="utf-8") as f:
        tags = TagList(json.load(f))
    crosswalk_rows = cw.load_crosswalk(args.crosswalk)
    edits = read_edits(args.workbook)

    print(f"[INFO] DRY RUN. Nothing is sent to OLM.\n[INFO] Backup (raw export): {photos_path}")
    print(f"[INFO] Edited photos in workbook: {len(edits)}\n")

    plan, blocked = [], 0
    for photo_id in sorted(edits):
        photo, problems = photos.get(photo_id), []
        if photo is None:
            print(f"Photo {photo_id}: BLOCKED, not in the raw export (deleted, or export is older than the edit)\n")
            blocked += 1
            continue
        raw_tags = photo.get("new_tags") or []
        before = [to_payload(t) for t in raw_tags]
        if not matches_summary(before, photo):
            problems.append("rebuilt tags do not match OLM's summary, so a replay might not be exact")
        found = {t["id"] for t in raw_tags}
        for tag_id in edits[photo_id]:
            if tag_id not in found:
                problems.append(f"edited tag {tag_id} is no longer on the photo (re-export and rebuild the workbook)")
        after = [apply_change(p, edits[photo_id].get(t["id"], {}), tags, problems) if t["id"] in edits[photo_id] else p
                 for t, p in zip(raw_tags, before)]
        for t, b, a in zip(raw_tags, before, after):
            if b["quantity"] != a["quantity"]:
                problems.append("quantity would change")
            if b["picked_up"] != a["picked_up"] and CHANGE_PICKED not in edits[photo_id].get(t["id"], {}):
                problems.append("picked_up would change without a Change_to_Picked_Up request")

        print(f"Photo {photo_id}  ({str(photo.get('datetime'))[:10]}, verified={photo.get('verified')})")
        for t, b, a in zip(raw_tags, before, after):
            if a == b:
                print(f"    same    {describe(b, tags)}")
            else:
                print(f"    before  {describe(b, tags)}\n    after   {describe(a, tags)}"
                      f"\n            -> {layer_for(a, tags, crosswalk_rows)}")
        for p in problems:
            print(f"    PROBLEM {p}")
        print()
        if problems:
            blocked += 1
        plan.append({"photo_id": photo_id, "ok": not problems, "problems": problems,
                     "verified_before": photo.get("verified"),
                     "backup_tags": raw_tags,
                     "backup_payload": {"photo_id": photo_id, "tags": before},
                     "target_payload": {"photo_id": photo_id, "tags": after}})

    os.makedirs(args.out_dir, exist_ok=True)
    out_path = os.path.join(args.out_dir, f"{args.label}_dryrun_{datetime.now():%Y%m%d-%H%M%S}.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump({"label": args.label, "workbook": args.workbook, "raw_export": photos_path, "photos": plan}, f, indent=1)
    os.chmod(out_path, 0o600)

    ready = sum(1 for p in plan if p["ok"])
    print(f"[SUMMARY] {ready} photos ready, {blocked} blocked. Plan saved to {out_path}")
    return 0 if blocked == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
